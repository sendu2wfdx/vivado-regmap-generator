#!/usr/bin/env python3
"""Convert the legacy 32-rows-per-register workbook to gen_regmap JSON."""
import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


def merged_end_row(sheet, row, column):
    for area in sheet.merged_cells.ranges:
        if area.min_col <= column <= area.max_col and area.min_row <= row <= area.max_row:
            return area.max_row
    return row


def clean_description(value):
    if not value:
        return ""
    text = str(value).strip()
    # The legacy file contains already-damaged mojibake in several descriptions.
    return "" if "�" in text else text


def convert(source: Path):
    values = load_workbook(source, data_only=True)
    blocks = []
    for sheet in values.worksheets:
        registers = []
        starts = [row for row in range(2, sheet.max_row + 1) if sheet.cell(row, 1).value]
        for start in starts:
            name = str(sheet.cell(start, 1).value)
            offset = int(str(sheet.cell(start, 2).value), 16)
            reset = int(str(sheet.cell(start, 3).value), 16)
            fields = []
            for row in range(start, min(start + 32, sheet.max_row + 1)):
                field_name = sheet.cell(row, 7).value
                if not field_name or str(field_name).upper() == "RESERVED":
                    continue
                end = merged_end_row(sheet, row, 7)
                field = {
                    "name": str(field_name),
                    "msb": int(sheet.cell(row, 5).value),
                    "lsb": int(sheet.cell(end, 5).value),
                    "access": str(sheet.cell(row, 8).value or "RO"),
                }
                description = clean_description(sheet.cell(row, 9).value)
                if description:
                    field["description"] = description
                fields.append(field)
            registers.append({
                "name": name,
                "offset": f"0x{offset:04X}",
                "reset": f"0x{reset:08X}",
                "fields": fields,
            })
        blocks.append({"name": sheet.title, "registers": registers})
    return {"prefix": "ADS4229", "blocks": blocks}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    args.output.write_text(json.dumps(convert(args.source), ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
