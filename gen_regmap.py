#!/usr/bin/env python3
"""Generate an Excel register map, Markdown table and C header from JSON.

Usage:
  python gen_regmap.py register_map.json regmap.xlsx \
      --markdown regmap.md --header regmap.h
"""

from __future__ import annotations

import argparse
import json
import re
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HEADERS = [
    "Register Name", "Offset (hex)", "Reg Value\n(hex)", "Byte Value\n(hex)",
    "Bits", "Bit Value\n(hex)", "Field Name", "Memory Access", "Description",
]
ACCESS = {"RO", "RW", "WO", "W1S", "W1C", "W1T", "RC", "RSVD"}


def number(value) -> int:
    if isinstance(value, int):
        return value
    return int(str(value), 0)


def normalize_register(reg: dict) -> dict:
    result = dict(reg)
    result["offset"] = number(reg["offset"])
    result["reset"] = number(reg.get("reset", 0))
    if not 0 <= result["reset"] <= 0xFFFFFFFF:
        raise ValueError(f"{reg['name']}: reset value is outside 32 bits")

    occupied = [False] * 32
    fields = []
    for src in reg.get("fields", []):
        field = dict(src)
        field["msb"] = int(field["msb"])
        field["lsb"] = int(field.get("lsb", field["msb"]))
        field["access"] = field.get("access", "RW").upper()
        if field["access"] not in ACCESS:
            raise ValueError(f"{reg['name']}.{field['name']}: invalid access")
        if not 0 <= field["lsb"] <= field["msb"] <= 31:
            raise ValueError(f"{reg['name']}.{field['name']}: invalid bit range")
        for bit in range(field["lsb"], field["msb"] + 1):
            if occupied[bit]:
                raise ValueError(f"{reg['name']}: overlapping field at bit {bit}")
            occupied[bit] = True
        fields.append(field)

    # Add reserved ranges automatically so every register covers all 32 bits.
    bit = 31
    while bit >= 0:
        if occupied[bit]:
            bit -= 1
            continue
        msb = bit
        while bit >= 0 and not occupied[bit]:
            occupied[bit] = True
            bit -= 1
        fields.append({
            "name": "RESERVED", "msb": msb, "lsb": bit + 1,
            "access": "RO", "description": "Reserved; keep at reset value.",
        })
    result["fields"] = sorted(fields, key=lambda f: f["msb"], reverse=True)
    return result


def load_schema(path: Path) -> dict:
    data = json.loads(path.read_text(encoding="utf-8"))
    if not data.get("blocks"):
        raise ValueError("schema must contain at least one block")
    for block in data["blocks"]:
        seen = set()
        regs = []
        for reg in block.get("registers", []):
            item = normalize_register(reg)
            if item["offset"] in seen:
                raise ValueError(f"{block['name']}: duplicate offset 0x{item['offset']:X}")
            seen.add(item["offset"])
            regs.append(item)
        block["registers"] = sorted(regs, key=lambda r: r["offset"])
    return data


def field_for_bit(reg: dict, bit: int) -> dict:
    return next(f for f in reg["fields"] if f["lsb"] <= bit <= f["msb"])


def make_workbook(schema: dict) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="B7C9DA")
    medium = Side(style="medium", color="5B7894")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    register_fills = [PatternFill("solid", fgColor="EAF2F8"), PatternFill("solid", fgColor="F7FAFC")]
    reserved_fill = PatternFill("solid", fgColor="E7E6E6")
    access_fills = {
        "RO": PatternFill("solid", fgColor="D9EAF7"),
        "RW": PatternFill("solid", fgColor="E2F0D9"),
        "WO": PatternFill("solid", fgColor="FFF2CC"),
        "W1S": PatternFill("solid", fgColor="FCE4D6"),
        "W1C": PatternFill("solid", fgColor="FCE4D6"),
        "W1T": PatternFill("solid", fgColor="FCE4D6"),
        "RC": PatternFill("solid", fgColor="DDEBF7"),
    }

    for block in schema["blocks"]:
        ws = wb.create_sheet(re.sub(r"[\\/*?:\[\]]", "_", block["name"])[:31])
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A2"
        ws.append(HEADERS)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=medium)
        ws.row_dimensions[1].height = 32

        row = 2
        for reg_index, reg in enumerate(block["registers"]):
            first, last = row, row + 31
            ws.merge_cells(start_row=first, start_column=1, end_row=last, end_column=1)
            ws.merge_cells(start_row=first, start_column=2, end_row=last, end_column=2)
            ws.merge_cells(start_row=first, start_column=3, end_row=last, end_column=3)
            ws.cell(first, 1, reg["name"])
            ws.cell(first, 2, f'=DEC2HEX({reg["offset"]},4)&\"h\"')
            ws.cell(first, 3, f'=LEFT(D{first},2)&LEFT(D{first+8},2)&LEFT(D{first+16},2)&LEFT(D{first+24},2)&\"h\"')

            for byte in range(4):
                br = first + byte * 8
                ws.merge_cells(start_row=br, start_column=4, end_row=br + 7, end_column=4)
                terms = "+".join(f"{1 << (7-i)}*F{br+i}" for i in range(8))
                ws.cell(br, 4, f'=DEC2HEX({terms},2)&\"h\"')

            for bit in range(31, -1, -1):
                rr = first + (31 - bit)
                ws.cell(rr, 5, bit)
                ws.cell(rr, 6, (reg["reset"] >> bit) & 1)

            for field in reg["fields"]:
                fr = first + (31 - field["msb"])
                lr = first + (31 - field["lsb"])
                for col in (7, 8, 9):
                    if lr > fr:
                        ws.merge_cells(start_row=fr, start_column=col, end_row=lr, end_column=col)
                ws.cell(fr, 7, field["name"])
                ws.cell(fr, 8, field["access"])
                ws.cell(fr, 9, field.get("description", ""))
                fill = reserved_fill if field["name"] == "RESERVED" else access_fills.get(field["access"])
                if fill:
                    for cells in ws.iter_rows(min_row=fr, max_row=lr, min_col=7, max_col=9):
                        for cell in cells:
                            cell.fill = fill

            base_fill = register_fills[reg_index % 2]
            for cells in ws.iter_rows(min_row=first, max_row=last, min_col=1, max_col=9):
                for cell in cells:
                    if cell.fill.fill_type is None:
                        cell.fill = base_fill
                    cell.font = Font(name="Arial", size=9)
                    cell.alignment = Alignment(
                        horizontal="left" if cell.column in (1, 7, 9) else "center",
                        vertical="center", wrap_text=cell.column == 9,
                    )
                    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            for col in range(1, 10):
                ws.cell(first, col).border = Border(left=thin, right=thin, top=medium, bottom=thin)
                ws.cell(last, col).border = Border(left=thin, right=thin, top=thin, bottom=medium)
            row += 32

        widths = [23, 14, 16, 15, 8, 14, 24, 15, 52]
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.auto_filter.ref = f"A1:I{max(1, ws.max_row)}"
        ws.print_title_rows = "1:1"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_margins.left = ws.page_margins.right = 0.25

    return wb


def patch_formula_cache(path: Path, schema: dict) -> None:
    """Store initial formula results so non-calculating XLSX viewers show them."""
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ET.register_namespace("", ns)
    cache_by_sheet = []
    for block in schema["blocks"]:
        cache = {}
        row = 2
        for reg in block["registers"]:
            cache[f"B{row}"] = f"{reg['offset']:04X}h"
            cache[f"C{row}"] = f"{reg['reset']:08X}h"
            for byte in range(4):
                value = (reg["reset"] >> (24 - 8 * byte)) & 0xFF
                cache[f"D{row + byte * 8}"] = f"{value:02X}h"
            row += 32
        cache_by_sheet.append(cache)

    with tempfile.NamedTemporaryFile(suffix=".xlsx", dir=path.parent, delete=False) as tmp:
        temp_path = Path(tmp.name)
    try:
        with zipfile.ZipFile(path, "r") as src, zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as dst:
            for info in src.infolist():
                data = src.read(info.filename)
                match = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", info.filename)
                if match:
                    index = int(match.group(1)) - 1
                    if index < len(cache_by_sheet):
                        root = ET.fromstring(data)
                        for cell in root.findall(f".//{{{ns}}}c"):
                            value = cache_by_sheet[index].get(cell.get("r"))
                            if value is None:
                                continue
                            cell.set("t", "str")
                            vnode = cell.find(f"{{{ns}}}v")
                            if vnode is None:
                                vnode = ET.SubElement(cell, f"{{{ns}}}v")
                            vnode.text = value
                        data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                dst.writestr(info, data)
        temp_path.replace(path)
    finally:
        temp_path.unlink(missing_ok=True)


def write_markdown(schema: dict, path: Path) -> None:
    lines = ["# Register Map", ""]
    for block in schema["blocks"]:
        lines += [f"## {block['name']}", "", "| Register | Offset | Bits | Field | Access | Reset | Description |", "|---|---:|---:|---|---|---:|---|"]
        for reg in block["registers"]:
            for field in reg["fields"]:
                bits = str(field["msb"]) if field["msb"] == field["lsb"] else f"{field['msb']}:{field['lsb']}"
                mask = (1 << (field["msb"] - field["lsb"] + 1)) - 1
                reset = (reg["reset"] >> field["lsb"]) & mask
                lines.append(f"| {reg['name']} | `0x{reg['offset']:04X}` | {bits} | {field['name']} | {field['access']} | `0x{reset:X}` | {field.get('description','')} |")
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def write_header(schema: dict, path: Path) -> None:
    guard = re.sub(r"\W+", "_", path.name.upper()) + "_"
    lines = [f"#ifndef {guard}", f"#define {guard}", ""]
    prefix = re.sub(r"\W+", "_", schema.get("prefix", "REGMAP").upper())
    for block in schema["blocks"]:
        lines.append(f"/* {block['name']} */")
        for reg in block["registers"]:
            rn = re.sub(r"\W+", "_", reg["name"].upper())
            lines.append(f"#define {prefix}_{rn}_OFFSET 0x{reg['offset']:04X}u")
            for field in reg["fields"]:
                if field["name"] == "RESERVED":
                    continue
                fn = re.sub(r"\W+", "_", field["name"].upper())
                width = field["msb"] - field["lsb"] + 1
                mask = ((1 << width) - 1) << field["lsb"]
                lines.append(f"#define {prefix}_{rn}_{fn}_SHIFT {field['lsb']}u")
                lines.append(f"#define {prefix}_{rn}_{fn}_MASK  0x{mask:08X}u")
            lines.append("")
    lines += [f"#endif /* {guard} */", ""]
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("schema", type=Path)
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("--markdown", type=Path)
    parser.add_argument("--header", type=Path)
    args = parser.parse_args()
    schema = load_schema(args.schema)
    args.xlsx.parent.mkdir(parents=True, exist_ok=True)
    workbook = make_workbook(schema)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(args.xlsx)
    patch_formula_cache(args.xlsx, schema)
    if args.markdown:
        write_markdown(schema, args.markdown)
    if args.header:
        write_header(schema, args.header)
    print(f"generated: {args.xlsx}")


if __name__ == "__main__":
    main()
